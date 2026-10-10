from io import BytesIO

from docx import Document
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader

from app.config import get_settings


async def extract_text(file: UploadFile) -> str:
    filename = (file.filename or "").lower()
    if not filename.endswith((".pdf", ".docx")):
        raise HTTPException(status_code=400, detail="Upload a PDF or DOCX file.")

    max_upload_size_bytes = get_settings().max_file_size_mb * 1024 * 1024
    content = await file.read(max_upload_size_bytes + 1)
    if len(content) > max_upload_size_bytes:
        raise HTTPException(
            status_code=413,
            detail=f"Resume files must be {get_settings().max_file_size_mb} MB or smaller.",
        )

    try:
        if filename.endswith(".pdf"):
            pages = PdfReader(BytesIO(content)).pages
            text = "\n".join(page.extract_text() or "" for page in pages)
        else:
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            text = "\n".join([*paragraphs, *table_cells])
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="The uploaded document could not be read.") from exc

    if len(text.strip()) < 50:
        raise HTTPException(status_code=400, detail="The resume must contain at least 50 characters of readable text.")
    return text.strip()